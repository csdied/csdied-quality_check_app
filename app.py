"""
Quality Check - Web app para auditoria de qualidade de respostas de pesquisa.
Versão otimizada para baixo consumo de memória (plano free 512MB).
Lê apenas as colunas necessárias do xlsx, não todas as 1948.
"""
import gc
import os
import secrets
from pathlib import Path

import pandas as pd
from flask import Flask, request, render_template, send_file, jsonify
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font

from quality import analyze

app = Flask(__name__)
app.secret_key = secrets.token_hex(16)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # 50 MB

WORK_DIR = Path("/tmp/qc_work")
WORK_DIR.mkdir(exist_ok=True)

# Colunas necessárias para análise. Ler só essas economiza memória.
NEEDED_COL_INDICES = [
    0,   # ID de respuesta
    2,   # Dirección IP
    5,   # Tiempo necesario para completar (segundos)
    63, 64, 65, 66, 67, 68,
    69, 70, 71, 72, 73, 74,
]


def read_only_needed_columns(filepath):
    """
    Lê só as colunas que precisamos pra análise — em vez de carregar
    todas as 1948 colunas. Reduz o uso de memória de ~400MB pra ~30MB.
    """
    df = pd.read_excel(
        filepath,
        sheet_name=0,
        usecols=NEEDED_COL_INDICES,
        engine='openpyxl',
    )
    return df


def make_lean_output(results_df, output_path):
    """
    Gera xlsx enxuto: só o resultado da auditoria.
    Cruzamento com dados originais é feito pelo ID na plataforma.
    """
    flag_cols = [
        'classificacao', 'score_suspeita',
        'flag_ip_duplicado', 'flag_tempo_curto', 'flag_marca_oleo',
        'flag_gibberish', 'flag_lugar_nao_marca', 'flag_so_simbolos',
        'flag_muitas_marcas_raras', 'flag_copy_paste',
        'motivos',
    ]
    main_cols = [
        'ID de respuesta', 'Dirección IP', 'Tiempo (min)',
    ] + flag_cols

    # Labels traduzidos para o relatório em espanhol
    flag_labels_es = {
        'flag_ip_duplicado': 'IP duplicada',
        'flag_tempo_curto': 'Tiempo muy corto',
        'flag_marca_oleo': 'Marca de aceite/automotriz',
        'flag_gibberish': 'Texto sin sentido',
        'flag_lugar_nao_marca': 'Lugar/institución como marca',
        'flag_so_simbolos': 'Sólo símbolos',
        'flag_muitas_marcas_raras': '3+ marcas raras',
        'flag_copy_paste': 'Copy-paste entre campos',
    }

    summary_rows = [['Total de encuestados', len(results_df)], ['', '']]
    summary_rows.append(['Clasificación', 'Cantidad'])
    for cls, n in results_df['classificacao'].value_counts().items():
        summary_rows.append([cls, int(n)])
    summary_rows.append(['', ''])
    summary_rows.append(['Detalle por bandera:', ''])
    for col in flag_cols:
        if col.startswith('flag_'):
            n = int(results_df[col].sum())
            summary_rows.append([flag_labels_es.get(col, col), n])
    df_summary = pd.DataFrame(summary_rows, columns=['Métrica', 'Valor'])

    df_suspects = results_df[results_df['score_suspeita'] >= 20].sort_values(
        'score_suspeita', ascending=False
    )[main_cols]

    df_full = results_df[main_cols].sort_values(
        'score_suspeita', ascending=False
    )

    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df_summary.to_excel(writer, sheet_name='Resumen', index=False)
        df_suspects.to_excel(writer, sheet_name='Revisar', index=False)
        df_full.to_excel(writer, sheet_name='Todos los IDs', index=False)

    wb = load_workbook(output_path)
    for sheet_name in ['Revisar', 'Todos los IDs']:
        ws = wb[sheet_name]
        header = [c.value for c in ws[1]]
        if 'classificacao' in header:
            cls_col = header.index('classificacao') + 1
            red_fill = PatternFill('solid', start_color='FFCDD2')
            yellow_fill = PatternFill('solid', start_color='FFF9C4')
            for row in range(2, ws.max_row + 1):
                val = ws.cell(row=row, column=cls_col).value
                if val and 'Alta suspeita' in str(val):
                    for col in range(1, ws.max_column + 1):
                        ws.cell(row=row, column=col).fill = red_fill
                elif val and 'Revisar' in str(val):
                    for col in range(1, ws.max_column + 1):
                        ws.cell(row=row, column=col).fill = yellow_fill

    ws_sum = wb['Resumen']
    ws_sum['A1'].font = Font(bold=True, size=12)
    ws_sum['B1'].font = Font(bold=True, size=12)
    ws_sum.column_dimensions['A'].width = 35
    ws_sum.column_dimensions['B'].width = 15
    wb.save(output_path)


@app.route('/')
def index():
    return render_template('index.html')


@app.route('/upload', methods=['POST'])
def upload():
    if 'file' not in request.files:
        return jsonify({'error': 'Nenhum arquivo enviado'}), 400
    f = request.files['file']
    if not f.filename:
        return jsonify({'error': 'Arquivo vazio'}), 400
    if not f.filename.lower().endswith(('.xlsx', '.xls')):
        return jsonify({'error': 'Por favor envie um arquivo .xlsx'}), 400
    try:
        time_threshold = int(request.form.get('time_threshold', 8))
    except ValueError:
        time_threshold = 8

    job_id = secrets.token_hex(8)
    input_path = WORK_DIR / f"{job_id}_input.xlsx"
    output_path = WORK_DIR / f"{job_id}_output.xlsx"
    f.save(input_path)

    try:
        df = read_only_needed_columns(input_path)
        # No df reduzido, marcas estão em índices 3 a 14
        brand_indices = list(range(3, 15))

        results, meta = analyze(
            df, time_threshold_min=time_threshold,
            brand_col_indices=brand_indices,
        )
        del df
        gc.collect()

        make_lean_output(results, output_path)

        try:
            input_path.unlink()
        except Exception:
            pass

        summary = {
            'job_id': job_id,
            'total': len(results),
            'classificacoes': {
                str(k): int(v) for k, v in
                results['classificacao'].value_counts().to_dict().items()
            },
            'flags': {
                col: int(results[col].sum())
                for col in results.columns if col.startswith('flag_')
            },
            'whitelist_size': len(meta['whitelist']),
            'whitelist_sample': sorted(meta['whitelist']),
            'tempo_minimo_min': time_threshold,
            # Manda TODOS os respondentes (ordenados por suspeita)
            'respondentes': results.sort_values(
                'score_suspeita', ascending=False
            )[
                ['ID de respuesta', 'Dirección IP', 'Tiempo (min)',
                 'classificacao', 'score_suspeita', 'motivos']
            ].fillna('').to_dict('records'),
        }

        del results
        gc.collect()
        return jsonify(summary)

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({
            'error': f'Erro ao processar: {type(e).__name__}: {e}'
        }), 500


@app.route('/download/<job_id>')
def download(job_id):
    if not all(c in '0123456789abcdef' for c in job_id) or len(job_id) != 16:
        return "Invalid job ID", 400
    output_path = WORK_DIR / f"{job_id}_output.xlsx"
    if not output_path.exists():
        return "Arquivo expirado ou não encontrado", 404
    return send_file(
        output_path,
        as_attachment=True,
        download_name='auditoria_calidad.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@app.route('/health')
def health():
    return 'ok'


if __name__ == '__main__':
    port = int(os.environ.get('PORT', 5000))
    debug = os.environ.get('FLASK_DEBUG', '').lower() == 'true'
    app.run(debug=debug, host='0.0.0.0', port=port)
